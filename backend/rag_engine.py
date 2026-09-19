import os
import json
import networkx as nx
import pandas as pd
import docx
from pypdf import PdfReader

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import ChatOpenAI
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.messages import SystemMessage, HumanMessage, AIMessage
from dotenv import load_dotenv

load_dotenv()

# Enable streaming directly on the LLM
llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0, streaming=True)

# Hugging Face Embeddings
embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

vector_store = None
knowledge_graph = nx.Graph()

def extract_text_from_file(file_path: str) -> str:
    """Extracts raw text depending on file extension."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif ext in [".docx", ".doc"]:
        doc = docx.Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text])

    elif ext in [".txt", ".csv"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    elif ext in [".xlsx", ".xls"]:
        excel_data = pd.read_excel(file_path, sheet_name=None)
        sheets_text = []
        for sheet_name, df in excel_data.items():
            sheets_text.append(f"--- Sheet: {sheet_name} ---\n{df.to_string(index=False)}")
        return "\n\n".join(sheets_text)

    else:
        raise ValueError(f"Unsupported file format: {ext}")

def process_document(file_path: str):
    global vector_store, knowledge_graph
    
    # Extract text using the new handler
    text = extract_text_from_file(file_path)
    
    if not text.strip():
        raise ValueError("The uploaded file does not contain readable text.")
    
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
    chunks = splitter.split_text(text)
    docs = [Document(page_content=chunk) for chunk in chunks]
    
    if vector_store is None:
        vector_store = FAISS.from_documents(docs, embeddings)
    else:
        vector_store.add_documents(docs)
        
    for chunk in chunks:
        extract_and_add_triples(chunk)

def extract_and_add_triples(text_chunk: str):
    prompt = f"""
    Extract relationships from the text below. 
    Return ONLY a JSON array of objects with keys: "subject", "relation", "object".
    Text: {text_chunk}
    """
    try:
        # Using invoke instead of the deprecated predict
        response = llm.invoke(prompt)
        triples = json.loads(response.content)
        
        for t in triples:
            sub = t.get("subject", "").strip().lower()
            obj = t.get("object", "").strip().lower()
            rel = t.get("relation", "").strip().lower()
            if sub and obj and rel:
                knowledge_graph.add_edge(sub, obj, relation=rel)
    except Exception as e:
        print(f"Graph extraction failed for chunk: {e}")

async def ask_question_stream(query: str, history: list):
    if vector_store is None:
        yield "Please upload a document first."
        return

    # 1. Retrieve from Vector DB
    vector_results = vector_store.similarity_search(query, k=3)
    vector_context = "\n".join([doc.page_content for doc in vector_results])
    
    # 2. Retrieve from Knowledge Graph
    entities_prompt = f"Extract the main subjects/entities from this query as a comma-separated list: {query}"
    entities_response = await llm.ainvoke(entities_prompt)
    query_entities = [e.strip().lower() for e in entities_response.content.split(",")]
    
    graph_context = []
    for entity in query_entities:
        if entity in knowledge_graph:
            for neighbor in knowledge_graph.neighbors(entity):
                rel = knowledge_graph.edges[entity, neighbor]['relation']
                graph_context.append(f"{entity} --[{rel}]--> {neighbor}")
                
    graph_context_str = "\n".join(graph_context)
    
    # 3. Combine contexts & history
    system_prompt = f"""
    You are an AI assistant helping with a document. Answer based on the context provided.
    
    Vector Context (Semantic matches):
    {vector_context}
    
    Knowledge Graph Context (Factual relationships):
    {graph_context_str}
    """
    
    messages = [SystemMessage(content=system_prompt)]
    
    # Inject conversation history
    for msg in history:
        if msg["role"] == "user":
            messages.append(HumanMessage(content=msg["content"]))
        elif msg["role"] == "assistant":
            messages.append(AIMessage(content=msg["content"]))
            
    # Append the new query
    messages.append(HumanMessage(content=query))
    
    # 4. Stream response back to FastAPI
    async for chunk in llm.astream(messages):
        if chunk.content:
            yield chunk.content